"""
Backend Logic - Ultimate Duplicate Finder
All core functionality extracted from the tkinter app, now runs as Eel backend.
NOTE: @eel.expose functions must be module-level, so we use a global instance pattern.
"""

import os
import shutil
import logging
import sqlite3
import threading
import hashlib
import time
import subprocess
import sys
import base64
import io
import tkinter as tk
from tkinter import filedialog

from PIL import Image
import imagehash
from thefuzz import fuzz
from send2trash import send2trash

try:
    import cv2
    OPENCV_AVAILABLE = True
except ImportError:
    OPENCV_AVAILABLE = False

try:
    import docx
    import openpyxl
    OFFICE_AVAILABLE = True
except ImportError:
    OFFICE_AVAILABLE = False

import eel

logging.basicConfig(
    filename='duplicate_finder.log',
    level=logging.INFO,
    format='%(asctime)s - %(message)s'
)


class _Backend:
    """Internal backend class. All state lives here."""
    def __init__(self):
        self.target_folders = []
        self.is_scanning = False
        self.scan_results = []  # list of dicts
        self.conn = None
        self.cursor = None
        self._init_db()

        # Hidden root for native dialogs (filedialog needs a Tk root)
        self._tk_root = tk.Tk()
        self._tk_root.withdraw()

    def _init_db(self):
        self.conn = sqlite3.connect('file_hash_cache.db', check_same_thread=False)
        self.cursor = self.conn.cursor()
        self.cursor.execute('''CREATE TABLE IF NOT EXISTS file_cache
                               (filepath TEXT PRIMARY KEY, mtime REAL,
                                hash_value TEXT, file_type TEXT)''')
        self.conn.commit()

    # -------- Hashing helpers (unchanged logic) --------

    def get_exact_hash(self, filepath):
        hasher = hashlib.md5()
        with open(filepath, 'rb') as f:
            for chunk in iter(lambda: f.read(65536), b""):
                hasher.update(chunk)
        return hasher.hexdigest()

    def get_docx_hash(self, filepath):
        doc = docx.Document(filepath)
        text = "\n".join([p.text for p in doc.paragraphs])
        return hashlib.md5(text.encode('utf-8')).hexdigest()

    def get_xlsx_hash(self, filepath):
        wb = openpyxl.load_workbook(filepath, data_only=True)
        full_text = []
        for s in wb.worksheets:
            for r in s.iter_rows(values_only=True):
                full_text.append(" ".join([str(c) for c in r if c is not None]))
        return hashlib.md5("\n".join(full_text).encode('utf-8')).hexdigest()

    def get_efficient_video_hash(self, filepath):
        if not OPENCV_AVAILABLE:
            return self.get_exact_hash(filepath)
        try:
            cap = cv2.VideoCapture(filepath)
            if not cap.isOpened():
                return "ERROR"
            total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
            if total_frames <= 0:
                cap.release()
                return "ERROR"

            checkpoints = [0.09, 0.18, 0.27, 0.36, 0.45,
                           0.54, 0.63, 0.72, 0.81, 0.90]
            hashes = []
            for cp in checkpoints:
                cap.set(cv2.CAP_PROP_POS_FRAMES, int(total_frames * cp))
                ret, frame = cap.read()
                if ret:
                    rgb_frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
                    h = str(imagehash.phash(Image.fromarray(rgb_frame)))
                    hashes.append(h)
            cap.release()
            return "".join(hashes) if hashes else "ERROR"
        except Exception:
            return "ERROR"

    def get_file_fingerprint(self, filepath, ext):
        mtime = os.path.getmtime(filepath)
        self.cursor.execute(
            "SELECT mtime, hash_value FROM file_cache WHERE filepath=?", (filepath,)
        )
        row = self.cursor.fetchone()
        if row and row[0] == mtime:
            return row[1]

        hash_val = "ERROR"
        try:
            if ext in ['jpg', 'jpeg', 'png', 'bmp']:
                hash_val = str(imagehash.phash(Image.open(filepath)))
            elif ext in ['mp4', 'mkv', 'avi', 'mov']:
                hash_val = self.get_efficient_video_hash(filepath)
            elif ext == 'docx' and OFFICE_AVAILABLE:
                hash_val = self.get_docx_hash(filepath)
            elif ext in ['xlsx', 'xls'] and OFFICE_AVAILABLE:
                hash_val = self.get_xlsx_hash(filepath)
            else:
                hash_val = self.get_exact_hash(filepath)

            if hash_val != "ERROR":
                self.cursor.execute(
                    "REPLACE INTO file_cache (filepath, mtime, hash_value, file_type) "
                    "VALUES (?, ?, ?, ?)",
                    (filepath, mtime, hash_val, ext),
                )
                self.conn.commit()
        except Exception:
            pass
        return hash_val

    # -------- Scan engine --------

    def _update_progress_eel(self, processed, total, start_time, ext):
        percent = (processed / total) * 100
        elapsed_time = time.time() - start_time
        if processed > 0:
            avg_time = elapsed_time / processed
            eta_seconds = int(avg_time * (total - processed))
            mins, secs = divmod(eta_seconds, 60)
            hours, mins = divmod(mins, 60)
            eta_str = (f"ETA: {hours:02d}:{mins:02d}:{secs:02d}"
                       if hours > 0 else f"ETA: {mins:02d}:{secs:02d}")
        else:
            eta_str = "ETA: Menghitung..."
        status_text = f"Memproses file {ext.upper()}... ({processed}/{total})"
        eel.updateProgress(percent, status_text, eta_str)

    def _add_result(self, file1, file2, status, ext):
        result = {
            "file1": file1,
            "file2": file2,
            "status": status,
            "ext": ext.upper(),
        }
        self.scan_results.append(result)
        eel.addResult(result)

    def _reset_scan_state(self):
        self.is_scanning = False
        eel.scanStateChanged(False)

    def _scan_process(self):
        files_by_ext = {}
        all_filepaths = []
        eel.updateProgress(0, "Menghitung total file...", "ETA: Menghitung...")

        for folder in self.target_folders:
            for root, _, files in os.walk(folder):
                for file in files:
                    filepath = os.path.join(root, file)
                    all_filepaths.append(filepath)
                    ext = filepath.split('.')[-1].lower()
                    if ext not in files_by_ext:
                        files_by_ext[ext] = []
                    files_by_ext[ext].append(filepath)

        total_files = len(all_filepaths)
        if total_files == 0:
            eel.updateProgress(100, "Tidak ada file ditemukan.", "ETA: -")
            self._reset_scan_state()
            return

        processed_files = 0
        start_time = time.time()

        for ext, file_list in files_by_ext.items():
            n = len(file_list)
            if n < 2:
                processed_files += n
                continue

            if ext in ['txt', 'csv']:
                for i in range(n):
                    processed_files += 1
                    self._update_progress_eel(processed_files, total_files,
                                              start_time, ext)
                    for j in range(i + 1, n):
                        try:
                            with open(file_list[i], 'r', encoding='utf-8',
                                      errors='ignore') as f1, \
                                 open(file_list[j], 'r', encoding='utf-8',
                                      errors='ignore') as f2:
                                if fuzz.token_sort_ratio(f1.read(), f2.read()) > 90:
                                    self._add_result(
                                        file_list[i], file_list[j],
                                        "Teks Mirip (>90%)", ext)
                        except Exception:
                            pass
                continue

            hash_dict = {}
            for filepath in file_list:
                processed_files += 1
                self._update_progress_eel(processed_files, total_files,
                                          start_time, ext)
                fingerprint = self.get_file_fingerprint(filepath, ext)
                if fingerprint and fingerprint != "ERROR":
                    if fingerprint in hash_dict:
                        original_path = hash_dict[fingerprint]
                        status = "Isi Identik"
                        if ext in ['jpg', 'jpeg', 'png', 'bmp']:
                            status = "Visual Mirip"
                        elif ext in ['mp4', 'mkv', 'avi', 'mov']:
                            status = "10-Frame Video Mirip"
                        elif ext in ['docx', 'xlsx']:
                            status = "Isi Dokumen Sama"
                        self._add_result(original_path, filepath, status, ext)
                    else:
                        hash_dict[fingerprint] = filepath

        eel.updateProgress(100, "Scan Selesai!", "ETA: Selesai")
        self._reset_scan_state()
        eel.scanComplete()


# ====================================================================
# Global singleton instance
# ====================================================================
_backend = _Backend()


# ====================================================================
# Eel-exposed functions (module-level, delegate to _backend)
# ====================================================================

@eel.expose
def add_folder():
    folder = filedialog.askdirectory(title="Pilih Folder untuk Ditambahkan")
    if folder and folder not in _backend.target_folders:
        _backend.target_folders.append(folder)
        return {"success": True, "folders": _backend.target_folders, "added": folder}
    return {"success": False, "folders": _backend.target_folders}


@eel.expose
def remove_folder(index):
    if 0 <= index < len(_backend.target_folders):
        removed = _backend.target_folders.pop(index)
        return {"success": True, "folders": _backend.target_folders, "removed": removed}
    return {"success": False, "folders": _backend.target_folders}


@eel.expose
def clear_folders():
    _backend.target_folders.clear()
    return {"success": True, "folders": _backend.target_folders}


@eel.expose
def get_folders():
    return _backend.target_folders


@eel.expose
def start_scan():
    if not _backend.target_folders:
        eel.showAlert("warning", "Peringatan",
                      "Tambahkan minimal 1 folder ke dalam daftar!")
        return
    if _backend.is_scanning:
        return

    _backend.is_scanning = True
    _backend.scan_results.clear()
    eel.scanStateChanged(True)

    thread = threading.Thread(target=_backend._scan_process, daemon=True)
    thread.start()


@eel.expose
def get_preview(filepath):
    """Generate a preview for a file. Returns dict with type/data keys."""
    if not os.path.exists(filepath):
        return {"type": "error", "message": "FILE TIDAK DITEMUKAN / SUDAH DIHAPUS"}

    ext = filepath.split('.')[-1].lower()
    try:
        if ext in ['jpg', 'jpeg', 'png', 'bmp', 'webp']:
            with Image.open(filepath) as img:
                img_copy = img.copy()
            img_copy.thumbnail((380, 280),
                               getattr(Image, 'Resampling', Image).LANCZOS)
            buffer = io.BytesIO()
            img_copy.save(buffer, format="PNG")
            img_b64 = base64.b64encode(buffer.getvalue()).decode('utf-8')
            return {"type": "image", "data": f"data:image/png;base64,{img_b64}"}

        elif ext in ['mp4', 'mkv', 'avi', 'mov', 'wmv', 'flv'] and OPENCV_AVAILABLE:
            cap = cv2.VideoCapture(filepath)
            if not cap.isOpened():
                return {"type": "error",
                        "message": "Gagal Buka Video: Codec Tidak Didukung"}
            cap.set(cv2.CAP_PROP_POS_FRAMES, 30)
            ret, frame = cap.read()
            if not ret:
                cap.set(cv2.CAP_PROP_POS_FRAMES, 0)
                ret, frame = cap.read()
            cap.release()
            if ret:
                rgb_frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
                img = Image.fromarray(rgb_frame)
                img.thumbnail((380, 280),
                              getattr(Image, 'Resampling', Image).LANCZOS)
                buffer = io.BytesIO()
                img.save(buffer, format="PNG")
                img_b64 = base64.b64encode(buffer.getvalue()).decode('utf-8')
                return {"type": "image", "data": f"data:image/png;base64,{img_b64}"}
            else:
                return {"type": "error",
                        "message": "Video Rusak atau Tidak Ada Visual"}

        elif ext in ['txt', 'csv']:
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                snippet = f.read(500)
            return {"type": "text", "data": snippet}

        else:
            filename = os.path.basename(filepath)
            filesize = os.path.getsize(filepath) / (1024 * 1024)
            return {
                "type": "info",
                "data": f"DOKUMEN: {ext.upper()}\n{filename}\nUkuran: {filesize:.2f} MB",
            }

    except Exception as e:
        return {"type": "error", "message": f"Error: {str(e)[:60]}"}


@eel.expose
def delete_file2(result_index):
    """Delete file2 of a specific result entry."""
    if 0 <= result_index < len(_backend.scan_results):
        file2 = _backend.scan_results[result_index]["file2"]
        safe_path = os.path.normpath(os.path.abspath(file2))
        try:
            if os.path.exists(safe_path):
                send2trash(safe_path)
            try:
                _backend.cursor.execute(
                    "DELETE FROM file_cache WHERE filepath=?", (file2,))
                _backend.conn.commit()
            except Exception as db_e:
                logging.error(
                    f"Gagal menghapus entri database untuk {file2}: {db_e}")
            return {"success": True, "file": file2}
        except Exception as e:
            return {"success": False, "error": str(e)}
    return {"success": False, "error": "Invalid index"}


@eel.expose
def move_file2(result_index, dest_folder):
    """Move file2 of a specific result entry to dest_folder."""
    if 0 <= result_index < len(_backend.scan_results):
        file2 = _backend.scan_results[result_index]["file2"]
        safe_path = os.path.normpath(os.path.abspath(file2))
        try:
            if os.path.exists(safe_path):
                shutil.move(safe_path, dest_folder)
            try:
                _backend.cursor.execute(
                    "DELETE FROM file_cache WHERE filepath=?", (file2,))
                _backend.conn.commit()
            except Exception as db_e:
                logging.error(
                    f"Gagal menghapus entri database untuk {file2}: {db_e}")
            return {"success": True, "file": file2, "dest": dest_folder}
        except Exception as e:
            return {"success": False, "error": str(e)}
    return {"success": False, "error": "Invalid index"}


@eel.expose
def open_file_location(result_index, file_index):
    """Open file location in explorer. file_index: 0=original, 1=duplicate."""
    key = f"file{file_index + 1}"
    if 0 <= result_index < len(_backend.scan_results) \
            and key in _backend.scan_results[result_index]:
        file_path = _backend.scan_results[result_index][key]
        try:
            if sys.platform == 'win32':
                subprocess.Popen(
                    f'explorer /select,"{os.path.normpath(file_path)}"')
            elif sys.platform == 'darwin':
                subprocess.Popen(['open', '-R', file_path])
            else:
                folder_path = os.path.dirname(file_path)
                subprocess.Popen(['xdg-open', folder_path])
            return {"success": True}
        except Exception as e:
            return {"success": False, "error": str(e)}
    return {"success": False, "error": "Invalid index"}


@eel.expose
def choose_destination():
    """Open native folder picker for move destination."""
    folder = filedialog.askdirectory(title="Pilih Folder Tujuan")
    return folder if folder else None

