import os
import shutil
import logging
import sqlite3
import threading
import hashlib
import time
import subprocess
import sys
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from PIL import Image, ImageTk
import imagehash
from thefuzz import fuzz
from send2trash import send2trash 

try:
    import cv2
    OPENCV_AVAILABLE = True
except ImportError:
    OPENCV_AVAILABLE = False

try:
    import docx
    import openpyxl
    OFFICE_AVAILABLE = True
except ImportError:
    OFFICE_AVAILABLE = False

logging.basicConfig(filename='duplicate_finder.log', level=logging.INFO, format='%(asctime)s - %(message)s')

class DuplicateFinderApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Ultimate Duplicate Finder (Full Pro & Safe Delete)")
        self.root.geometry("1300x850") 
        
        self.target_folders = []
        self.is_scanning = False
        
        self.thumb1_tk = None
        self.thumb2_tk = None
        
        self.init_db()
        self.setup_ui()

    def init_db(self):
        self.conn = sqlite3.connect('file_hash_cache.db', check_same_thread=False)
        self.cursor = self.conn.cursor()
        self.cursor.execute('''CREATE TABLE IF NOT EXISTS file_cache 
                               (filepath TEXT PRIMARY KEY, mtime REAL, hash_value TEXT, file_type TEXT)''')
        self.conn.commit()

    def setup_ui(self):
        # --- PANEL ATAS ---
        top_frame = tk.Frame(self.root, padx=10, pady=10)
        top_frame.pack(fill=tk.X)
        
        tk.Label(top_frame, text="Daftar Folder yang Akan Di-scan (Otomatis termasuk subfolder):", font=("Arial", 10, "bold")).pack(anchor="w")
        self.listbox_folders = tk.Listbox(top_frame, height=4, width=80)
        self.listbox_folders.pack(side=tk.LEFT, fill=tk.X, expand=True, pady=5)
        
        btn_frame = tk.Frame(top_frame)
        btn_frame.pack(side=tk.LEFT, padx=10)
        tk.Button(btn_frame, text="+ Tambah Folder", command=self.add_folder, width=15).pack(pady=2)
        tk.Button(btn_frame, text="- Hapus Terpilih", command=self.remove_folder, width=15).pack(pady=2)
        tk.Button(btn_frame, text="Bersihkan Semua", command=self.clear_folders, width=15).pack(pady=2)

        # --- PANEL TENGAH ---
        progress_frame = tk.Frame(self.root, padx=10, pady=5)
        progress_frame.pack(fill=tk.X)
        
        self.btn_scan = tk.Button(progress_frame, text="MULAI SCAN DUPLIKAT", command=self.start_scan_thread, bg="green", fg="white", font=("Arial", 10, "bold"))
        self.btn_scan.pack(side=tk.LEFT, padx=5)
        
        self.lbl_status = tk.Label(progress_frame, text="Status: Menunggu", fg="blue", width=40, anchor="w")
        self.lbl_status.pack(side=tk.LEFT, padx=5)
        
        self.lbl_eta = tk.Label(progress_frame, text="ETA: -", fg="red")
        self.lbl_eta.pack(side=tk.RIGHT, padx=5)

        self.progress_var = tk.DoubleVar()
        self.progress_bar = ttk.Progressbar(progress_frame, variable=self.progress_var, maximum=100)
        self.progress_bar.pack(side=tk.BOTTOM, fill=tk.X, pady=5, padx=5)

        # --- PANEL BAWAH (SPLIT VIEW) ---
        main_frame = tk.PanedWindow(self.root, orient=tk.HORIZONTAL)
        main_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
        # Kiri: Tabel + Scrollbar
        tree_frame = tk.Frame(main_frame)
        self.tree_scroll = ttk.Scrollbar(tree_frame, orient="vertical")
        self.tree = ttk.Treeview(tree_frame, columns=("File 1", "File 2", "Kemiripan", "Tipe"), 
                                 show='headings', yscrollcommand=self.tree_scroll.set)
        
        self.tree_scroll.config(command=self.tree.yview)
        self.tree_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self.tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        self.tree.heading("File 1", text="File Original")
        self.tree.heading("File 2", text="Terduga Duplikat")
        self.tree.heading("Kemiripan", text="Status")
        self.tree.heading("Tipe", text="Tipe")
        self.tree.column("Kemiripan", width=120)
        self.tree.column("Tipe", width=70)
        self.tree.bind('<<TreeviewSelect>>', self.on_tree_select)
        main_frame.add(tree_frame, stretch="always")
        
        # Kanan: Panel Aksi & Preview
        right_panel = tk.Frame(main_frame, width=450, padx=10)
        
        btn_action_frame = tk.Frame(right_panel)
        btn_action_frame.pack(fill=tk.X, pady=5)
        
        tk.Label(btn_action_frame, text="Tindakan Eksekusi:", font=("Arial", 10, "bold")).grid(row=0, column=0, columnspan=2, pady=2, sticky="w")
        self.btn_del = tk.Button(btn_action_frame, text="Hapus File 2 (Recycle Bin)", command=self.delete_selected, bg="#d9534f", fg="white", width=22)
        self.btn_del.grid(row=1, column=0, padx=2, pady=2)
        self.btn_mov = tk.Button(btn_action_frame, text="Pindahkan File 2", command=self.move_selected, width=20)
        self.btn_mov.grid(row=1, column=1, padx=2, pady=2)
        tk.Button(btn_action_frame, text="Buka Lokasi File 1", command=lambda: self.open_location(0), width=22).grid(row=2, column=0, padx=2, pady=2)
        tk.Button(btn_action_frame, text="Buka Lokasi File 2", command=lambda: self.open_location(1), width=20).grid(row=2, column=1, padx=2, pady=2)
        
        preview_label_frame = tk.LabelFrame(right_panel, text=" Big Live Preview ", font=("Arial", 9, "bold"), fg="purple", padx=5, pady=5)
        preview_label_frame.pack(fill=tk.BOTH, expand=True, pady=5)
        
        tk.Label(preview_label_frame, text="File 1 (Original):", font=("Arial", 9, "italic")).pack(anchor="w")
        self.canvas_p1 = tk.Label(preview_label_frame, text="[Pilih baris untuk melihat]", bg="#2b2b2b", fg="white", width=55, height=13)
        self.canvas_p1.pack(fill=tk.BOTH, expand=True, pady=2)
        
        tk.Label(preview_label_frame, text="File 2 (Duplikat):", font=("Arial", 9, "italic")).pack(anchor="w")
        self.canvas_p2 = tk.Label(preview_label_frame, text="[Pilih baris untuk melihat]", bg="#2b2b2b", fg="white", width=55, height=13)
        self.canvas_p2.pack(fill=tk.BOTH, expand=True, pady=2)
        
        main_frame.add(right_panel)

    def add_folder(self):
        folder = filedialog.askdirectory(title="Pilih Folder untuk Ditambahkan")
        if folder and folder not in self.target_folders:
            self.target_folders.append(folder)
            self.listbox_folders.insert(tk.END, folder)

    def remove_folder(self):
        selected = self.listbox_folders.curselection()
        if selected:
            idx = selected[0]
            self.listbox_folders.delete(idx)
            del self.target_folders[idx]

    def clear_folders(self):
        self.listbox_folders.delete(0, tk.END)
        self.target_folders.clear()

    def on_tree_select(self, event):
        selected = self.tree.selection()
        if not selected: return
        
        file1, file2, status, ext = self.tree.item(selected[0], 'values')
        ext = ext.lower()
        
        self.thumb1_tk = self.generate_preview_data(file1, ext, self.canvas_p1)
        self.thumb2_tk = self.generate_preview_data(file2, ext, self.canvas_p2)

    def generate_preview_data(self, filepath, ext, target_widget):
        if not os.path.exists(filepath):
            target_widget.config(image='', text="[FILE TIDAK DITEMUKAN / SUDAH DIHAPUS]", bg="#d9534f", fg="white")
            return None
            
        size = (380, 280) 
        
        try:
            if ext in ['jpg', 'jpeg', 'png', 'bmp', 'webp']:
                # PERBAIKAN: Gunakan context manager & copy() agar file lgsg ditutup oleh Windows
                with Image.open(filepath) as img:
                    img_copy = img.copy() 
                
                img_copy.thumbnail(size, getattr(Image, 'Resampling', Image).LANCZOS)
                img_tk = ImageTk.PhotoImage(img_copy)
                target_widget.config(image=img_tk, text="", bg="#2b2b2b")
                return img_tk
                
            elif ext in ['mp4', 'mkv', 'avi', 'mov', 'wmv', 'flv'] and OPENCV_AVAILABLE:
                cap = cv2.VideoCapture(filepath)
                if not cap.isOpened():
                    target_widget.config(image='', text="[Gagal Buka Video: Codec Tidak Didukung]", bg="#4a4a4a", fg="yellow")
                    return None
                
                cap.set(cv2.CAP_PROP_POS_FRAMES, 30)
                ret, frame = cap.read()
                
                if not ret:
                    cap.set(cv2.CAP_PROP_POS_FRAMES, 0)
                    ret, frame = cap.read()
                    
                cap.release() # VideoCapture sudah aman karena langsung di-release
                
                if ret:
                    rgb_frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
                    img = Image.fromarray(rgb_frame)
                    img.thumbnail(size, getattr(Image, 'Resampling', Image).LANCZOS)
                    img_tk = ImageTk.PhotoImage(img)
                    target_widget.config(image=img_tk, text="", bg="#2b2b2b")
                    return img_tk
                else:
                    target_widget.config(image='', text="[Video Rusak atau Tidak Ada Visual]", bg="#4a4a4a", fg="yellow")
                    
            elif ext in ['txt', 'csv']:
                with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                    snippet = f.read(250)
                target_widget.config(image='', text=f"PREVIEW TEKS:\n\n{snippet}...", bg="#ffffff", fg="black", anchor="nw", justify=tk.LEFT)
                
            else:
                filename = os.path.basename(filepath)
                filesize = os.path.getsize(filepath) / (1024 * 1024)
                target_widget.config(image='', text=f"DOKUMEN: {ext.upper()}\n\n{filename}\nUkuran: {filesize:.2f} MB", bg="#2b2b2b", fg="white", justify=tk.CENTER)
                
        except Exception as e:
            target_widget.config(image='', text=f"[ERROR MEMUAT PREVIEW]\n{str(e)[:40]}", bg="#d9534f", fg="white")
        return None

    def get_exact_hash(self, filepath):
        hasher = hashlib.md5()
        with open(filepath, 'rb') as f:
            for chunk in iter(lambda: f.read(65536), b""): hasher.update(chunk)
        return hasher.hexdigest()

    def get_docx_hash(self, filepath):
        doc = docx.Document(filepath)
        return hashlib.md5("\n".join([p.text for p in doc.paragraphs]).encode('utf-8')).hexdigest()

    def get_xlsx_hash(self, filepath):
        wb = openpyxl.load_workbook(filepath, data_only=True)
        full_text = [" ".join([str(c) for c in r if c is not None]) for s in wb.worksheets for r in s.iter_rows(values_only=True)]
        return hashlib.md5("\n".join(full_text).encode('utf-8')).hexdigest()

    def get_efficient_video_hash(self, filepath):
        if not OPENCV_AVAILABLE: return self.get_exact_hash(filepath)
        try:
            cap = cv2.VideoCapture(filepath)
            if not cap.isOpened(): return "ERROR"
            total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
            if total_frames <= 0: return "ERROR"

            checkpoints = [0.09, 0.18, 0.27, 0.36, 0.45, 0.54, 0.63, 0.72, 0.81, 0.90]
            hashes = []
            for cp in checkpoints:
                cap.set(cv2.CAP_PROP_POS_FRAMES, int(total_frames * cp))
                ret, frame = cap.read()
                if ret:
                    rgb_frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
                    hashes.append(str(imagehash.phash(Image.fromarray(rgb_frame))))
            cap.release()
            return "".join(hashes) if hashes else "ERROR"
        except: return "ERROR"

    def get_file_fingerprint(self, filepath, ext):
        mtime = os.path.getmtime(filepath)
        self.cursor.execute("SELECT mtime, hash_value FROM file_cache WHERE filepath=?", (filepath,))
        row = self.cursor.fetchone()
        if row and row[0] == mtime: return row[1]
            
        hash_val = "ERROR"
        try:
            if ext in ['jpg', 'jpeg', 'png', 'bmp']: hash_val = str(imagehash.phash(Image.open(filepath)))
            elif ext in ['mp4', 'mkv', 'avi', 'mov']: hash_val = self.get_efficient_video_hash(filepath)
            elif ext == 'docx' and OFFICE_AVAILABLE: hash_val = self.get_docx_hash(filepath)
            elif ext in ['xlsx', 'xls'] and OFFICE_AVAILABLE: hash_val = self.get_xlsx_hash(filepath)
            else: hash_val = self.get_exact_hash(filepath)
                
            if hash_val != "ERROR":
                self.cursor.execute("REPLACE INTO file_cache (filepath, mtime, hash_value, file_type) VALUES (?, ?, ?, ?)", (filepath, mtime, hash_val, ext))
                self.conn.commit()
        except: pass
        return hash_val

    def start_scan_thread(self):
        if not self.target_folders:
            messagebox.showwarning("Peringatan", "Tambahkan minimal 1 folder ke dalam daftar!")
            return
        if self.is_scanning: return
            
        self.is_scanning = True
        self.btn_scan.config(state=tk.DISABLED, bg="gray")
        self.tree.delete(*self.tree.get_children())
        
        thread = threading.Thread(target=self._scan_process, daemon=True)
        thread.start()

    def _scan_process(self):
        files_by_ext = {}
        all_filepaths = []
        self.update_ui("Menghitung total file (termasuk subfolder)...", 0, "ETA: Menghitung...")
        
        for folder in self.target_folders:
            for root, _, files in os.walk(folder):
                for file in files:
                    filepath = os.path.join(root, file)
                    all_filepaths.append(filepath)
                    ext = filepath.split('.')[-1].lower()
                    if ext not in files_by_ext: files_by_ext[ext] = []
                    files_by_ext[ext].append(filepath)

        total_files = len(all_filepaths)
        if total_files == 0:
            self.update_ui("Tidak ada file ditemukan.", 100, "ETA: -")
            self.reset_scan_state()
            return

        processed_files = 0
        start_time = time.time()

        for ext, file_list in files_by_ext.items():
            n = len(file_list)
            if n < 2:
                processed_files += n
                continue
                
            if ext in ['txt', 'csv']:
                for i in range(n):
                    processed_files += 1
                    self.update_progress(processed_files, total_files, start_time, ext)
                    for j in range(i + 1, n):
                        try:
                            with open(file_list[i], 'r', encoding='utf-8', errors='ignore') as f1, open(file_list[j], 'r', encoding='utf-8', errors='ignore') as f2:
                                if fuzz.token_sort_ratio(f1.read(), f2.read()) > 90:
                                    self.add_to_tree(file_list[i], file_list[j], "Teks Mirip (>90%)", ext)
                        except: pass
                continue

            hash_dict = {}
            for filepath in file_list:
                processed_files += 1
                self.update_progress(processed_files, total_files, start_time, ext)
                
                fingerprint = self.get_file_fingerprint(filepath, ext)
                if fingerprint and fingerprint != "ERROR":
                    if fingerprint in hash_dict:
                        original_path = hash_dict[fingerprint]
                        status = "Isi Identik"
                        if ext in ['jpg', 'jpeg', 'png', 'bmp']: status = "Visual Mirip"
                        elif ext in ['mp4', 'mkv', 'avi', 'mov']: status = "10-Frame Video Mirip"
                        elif ext in ['docx', 'xlsx']: status = "Isi Dokumen Sama"
                        
                        self.add_to_tree(original_path, filepath, status, ext)
                    else:
                        hash_dict[fingerprint] = filepath

        self.update_ui("Scan Selesai!", 100, "ETA: Selesai")
        self.reset_scan_state()
        messagebox.showinfo("Selesai", "Proses pemindaian selesai.")

    def update_progress(self, processed, total, start_time, ext):
        percent = (processed / total) * 100
        elapsed_time = time.time() - start_time
        if processed > 0:
            avg_time = elapsed_time / processed
            eta_seconds = int(avg_time * (total - processed))
            mins, secs = divmod(eta_seconds, 60)
            hours, mins = divmod(mins, 60)
            eta_str = f"ETA: {hours:02d}:{mins:02d}:{secs:02d}" if hours > 0 else f"ETA: {mins:02d}:{secs:02d}"
        else: eta_str = "ETA: Menghitung..."
            
        status_text = f"Memproses file {ext.upper()}... ({processed}/{total})"
        self.root.after(0, lambda: self.lbl_status.config(text=status_text))
        self.root.after(0, lambda: self.lbl_eta.config(text=eta_str))
        self.root.after(0, lambda: self.progress_var.set(percent))

    def update_ui(self, status, progress, eta):
        self.root.after(0, lambda: self.lbl_status.config(text=status))
        self.root.after(0, lambda: self.progress_var.set(progress))
        self.root.after(0, lambda: self.lbl_eta.config(text=eta))

    def reset_scan_state(self):
        self.is_scanning = False
        self.root.after(0, lambda: self.btn_scan.config(state=tk.NORMAL, bg="green"))

    def add_to_tree(self, file1, file2, status, ext):
        self.root.after(0, lambda: self.tree.insert("", tk.END, values=(file1, file2, status, ext.upper())))

    def delete_selected(self):
        selected_items = self.tree.selection()
        if not selected_items: return
        
        if messagebox.askyesno("Konfirmasi Penghapusan", "Apakah Anda yakin ingin memindahkan file terduga duplikat (File 2) ke Recycle Bin?"):
            for item in selected_items:
                file2 = self.tree.item(item, 'values')[1]
                
                # PERBAIKAN: Normalisasi path agar Windows tidak bingung & kebal error format path
                safe_path = os.path.normpath(os.path.abspath(file2))
                
                try:
                    if os.path.exists(safe_path): 
                        send2trash(safe_path)
                except Exception as e: 
                    messagebox.showerror("Error", f"Gagal menghapus file:\n{e}")
                    continue 
                
                self.tree.delete(item)
                
                try:
                    self.cursor.execute("DELETE FROM file_cache WHERE filepath=?", (file2,))
                    self.conn.commit()
                except Exception as db_e:
                    logging.error(f"Gagal menghapus entri database untuk {file2}: {db_e}")
                
                self.canvas_p1.config(image='', text="[Pilih baris untuk melihat]", bg="#2b2b2b")
                self.canvas_p2.config(image='', text="[Pilih baris untuk melihat]", bg="#2b2b2b")

    def move_selected(self):
        selected_items = self.tree.selection()
        if not selected_items: return
        dest_folder = filedialog.askdirectory(title="Pilih Folder Tujuan")
        if not dest_folder: return
        for item in selected_items:
            try:
                file2 = self.tree.item(item, 'values')[1]
                safe_path = os.path.normpath(os.path.abspath(file2))
                if os.path.exists(safe_path): shutil.move(safe_path, dest_folder)
                self.tree.delete(item)
                self.cursor.execute("DELETE FROM file_cache WHERE filepath=?", (file2,))
                self.conn.commit()
                self.canvas_p1.config(image='', text="[Pilih baris untuk melihat]", bg="#2b2b2b")
                self.canvas_p2.config(image='', text="[Pilih baris untuk melihat]", bg="#2b2b2b")
            except Exception as e: messagebox.showerror("Error", f"Gagal memindahkan:\n{e}")

    def open_location(self, file_index):
        selected = self.tree.selection()
        if not selected: return
        file_path = self.tree.item(selected[0], 'values')[file_index]
        try:
            if sys.platform == 'win32': subprocess.Popen(f'explorer /select,"{os.path.normpath(file_path)}"')
            elif sys.platform == 'darwin': subprocess.Popen(['open', '-R', file_path])
            else:
                folder_path = os.path.dirname(file_path)
                subprocess.Popen(['xdg-open', folder_path])
        except Exception as e: messagebox.showerror("Error", f"Gagal membuka lokasi:\n{e}")

if __name__ == "__main__":
    root = tk.Tk()
    app = DuplicateFinderApp(root)
    root.mainloop()